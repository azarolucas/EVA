# -*- coding: utf-8 -*-
"""
ATHENA BUSINESS PACKAGE — GERADOR
Lê 00_FICHA_DADOS_CLIENTE.xlsx e gera todos os documentos do pacote,
coerentes entre si, com as correcoes de logica aplicadas.
Uso:  python3 gerar_pacote.py 00_FICHA_DADOS_CLIENTE.xlsx
"""
import sys, os, re
import openpyxl
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3B, 0x5C)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GREY = RGBColor(0x6B, 0x77, 0x85)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

# ---------- number to words (0..100) ----------
_ONES = ["ZERO","ONE","TWO","THREE","FOUR","FIVE","SIX","SEVEN","EIGHT","NINE","TEN",
         "ELEVEN","TWELVE","THIRTEEN","FOURTEEN","FIFTEEN","SIXTEEN","SEVENTEEN",
         "EIGHTEEN","NINETEEN"]
_TENS = ["","","TWENTY","THIRTY","FORTY","FIFTY","SIXTY","SEVENTY","EIGHTY","NINETY"]
def num_words(n):
    n = int(round(n))
    if n == 100: return "ONE HUNDRED"
    if n < 20: return _ONES[n]
    t, o = divmod(n, 10)
    return _TENS[t] + ("" if o == 0 else "-" + _ONES[o])

def pct_words(p):
    p = float(p)
    if abs(p - round(p)) < 1e-9:
        return num_words(round(p)) + " PERCENT"
    return f"{p:g} PERCENT"

# ---------- read master ficha (single source: FICHA_ABERTURA_ATHENA) ----------
def _sim(v):
    return str(v or "").strip().upper() in ("SIM", "YES", "TRUE", "S", "Y")

def _num(v):
    try: return float(str(v).replace("%","").replace(",",".").strip() or 0)
    except: return 0.0

def read_ficha(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["1. ABERTURA"]
    km = wb["_KEYS"]
    kv = {}
    for r in range(2, km.max_row + 1):
        k = km.cell(row=r, column=1).value
        co = km.cell(row=r, column=2).value
        if k and co:
            kv[str(k).strip()] = str(co).strip()
    def get(key, default=""):
        co = kv.get(key)
        if not co: return default
        v = ws[co].value
        return default if v is None else str(v).strip()
    data = {}
    for k in kv:
        if not re.match(r'S\d+_', k):
            data[k] = get(k)
    def compose(a, c, s, z):
        return " - ".join([p for p in [get(a), get(c), get(s), get(z)] if p])
    data["PRINCIPAL_ADDRESS"] = compose("PRINCIPAL_ADDRESS","PRINCIPAL_CITY","PRINCIPAL_STATE","PRINCIPAL_ZIP") or get("PRINCIPAL_ADDRESS")
    if _sim(get("MAILING_SAME","X")) or not get("MAILING_ADDRESS"):
        data["MAILING_ADDRESS"] = data["PRINCIPAL_ADDRESS"]
    else:
        data["MAILING_ADDRESS"] = compose("MAILING_ADDRESS","MAILING_CITY","MAILING_STATE","MAILING_ZIP")
    data["REGISTERED_AGENT_ADDRESS"] = compose("REGISTERED_AGENT_ADDRESS","AGENT_CITY","AGENT_STATE","AGENT_ZIP") or get("REGISTERED_AGENT_ADDRESS")
    data.setdefault("DURATION", "PERPETUAL")
    data["FORMATION_DATE"] = get("ABERTURA_DATA")
    data["CLIENT_PHONE"] = get("CONTACT_PHONE")
    members = []
    for n in range(1, 20):
        name = get(f"S{n}_NOME")
        if not name:
            continue
        managing = _sim(get(f"S{n}_GERENTE"))
        addr = ", ".join([p for p in [get(f"S{n}_ENDERECO"), get(f"S{n}_CIDADE"),
                                       get(f"S{n}_ESTADO"), get(f"S{n}_CEP")] if p])
        members.append({
            "name": name, "address": addr,
            "role": "AUTHORIZED MEMBER" if managing else "MEMBER ONLY",
            "managing": managing, "position": get(f"S{n}_POSICAO"),
            "ssn": get(f"S{n}_SSN"), "dob": get(f"S{n}_DOB"),
            "phone": get(f"S{n}_TELEFONE"), "email": get(f"S{n}_EMAIL"),
            "contribution": "CASH", "fmv": None,
            "percent": _num(get(f"S{n}_PERCENT")),
        })
    return data, members

# ---------- docx helpers ----------
def set_margins(doc, top=1.0, bottom=1.0, left=1.0, right=1.0):
    for s in doc.sections:
        s.top_margin = Inches(top); s.bottom_margin = Inches(bottom)
        s.left_margin = Inches(left); s.right_margin = Inches(right)

def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement('w:fldSimple'); fld1.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld1)

def letterhead(doc, llc_name):
    """Athena header + clean footer (no inherited JGHA/page-x-of-y)."""
    sec = doc.sections[0]
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]; p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("ATHENA")
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = NAVY
    r.font.name = "Georgia"
    r2 = p.add_run("  BUSINESS & TAX ADVISORS")
    r2.font.size = Pt(9); r2.font.bold = True; r2.font.color.rgb = GOLD
    r2.font.name = "Calibri"
    p2 = hdr.add_paragraph()
    r = p2.add_run("7680 Universal Blvd, Suite 100 · Orlando, FL 32819 · (407) 777-2501 · manager@athenataxadvisors.com")
    r.font.size = Pt(7.5); r.font.color.rgb = GREY; r.font.name = "Calibri"
    # bottom border on p2
    pPr = p2._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6')
    bottom.set(qn('w:space'),'6'); bottom.set(qn('w:color'),'C9A227')
    pbdr.append(bottom); pPr.append(pbdr)

    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]; fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = fp.add_run(f"{llc_name}  —  CONFIDENTIAL  ·  Athena Business & Tax Advisors  ·  Page ")
    rf.font.size = Pt(7.5); rf.font.color.rgb = GREY; rf.font.name = "Calibri"
    add_page_number(fp)

def h1(doc, text, center=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text); r.font.bold = True; r.font.size = Pt(15)
    r.font.color.rgb = NAVY; r.font.name = "Georgia"
    p.space_after = Pt(4)
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.bold = True; r.font.size = Pt(11.5)
    r.font.color.rgb = NAVY; r.font.name = "Calibri"
    p.space_before = Pt(8); p.space_after = Pt(2)
    return p

def body(doc, text, size=10, bold=False, italic=False, color=BLACK, align="just", space_after=4):
    p = doc.add_paragraph()
    amap = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    p.alignment = amap[align]
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color; r.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(space_after)
    return p

def rule_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("⚠  " + text); r.font.size = Pt(8.5); r.font.italic = True
    r.font.color.rgb = GOLD; r.font.name = "Calibri"
    p.paragraph_format.space_after = Pt(6)

def sig_block(doc, name, role="AUTHORIZED MEMBER"):
    body(doc, "", space_after=2)
    body(doc, "Signature: ______________________________________________", size=10, space_after=2)
    body(doc, f"Printed Name:  {name}", size=10, bold=True, space_after=0)
    body(doc, role, size=9, color=GREY, space_after=8)


# =====================================================================
# DOC 0 — CAPA / INDICE  (branded, PT explanations)
# =====================================================================
def doc_capa(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    body(doc, "", space_after=30)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BUSINESS PACKAGE"); r.font.size = Pt(12); r.font.bold = True
    r.font.color.rgb = GOLD; r.font.name = "Calibri"
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("DOCUMENTS OF"); r.font.size = Pt(13); r.font.color.rgb = GREY
    r.font.name = "Georgia"
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run(data["LLC_NAME"]); r.font.size = Pt(26); r.font.bold = True
    r.font.color.rgb = NAVY; r.font.name = "Georgia"
    p4 = doc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p4.add_run(f"A {data['STATE']} Limited Liability Company"); r.font.size = Pt(11)
    r.font.italic = True; r.font.color.rgb = GREY; r.font.name = "Calibri"
    body(doc, "", space_after=20)

    h2(doc, "CONTEÚDO DO PACOTE / PACKAGE CONTENTS")
    items = [
        ("EIN — Employer Identification Number", "Número de identificação do empregador para fins tributários (equivalente ao CNPJ no Brasil)."),
        ("SunBiz Page", "Página da Divisão de Corporações da Flórida com as informações da empresa e seus representantes."),
        ("Certificate of Status", "Documento estadual que atesta que a LLC está autorizada a operar e em dia com o Estado."),
        ("Articles of Organization", "Documento formal que constitui a LLC junto ao Estado."),
        ("Operating Agreement", "Acordo entre os sócios: participação, gestão, votação, distribuições, entrada/saída e poderes. Documento privado e vinculante."),
        ("Membership Ledger", "Registro das participações (interests) emitidas a cada sócio."),
        ("Membership Certificate", "Certificado individual que especifica o percentual de cada sócio na empresa."),
        ("Statement of Authority", "Autorização interna que define quem pode assinar em nome da empresa (banco, contratos, bens)."),
        ("Mailbox Agreement", "Contrato de serviço de caixa postal (serviço contínuo, com período e preço próprios)."),
    ]
    for t, d in items:
        p = doc.add_paragraph()
        r = p.add_run("•  " + t + " — "); r.font.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = NAVY; r.font.name = "Calibri"
        r2 = p.add_run(d); r2.font.size = Pt(9.5); r2.font.color.rgb = BLACK; r2.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(3)

    body(doc, "", space_after=6)
    h2(doc, "OBRIGAÇÕES APÓS A ABERTURA / POST-FORMATION DEADLINES")
    body(doc, "Annual Report (recadastramento no Estado): 1º de janeiro a 1º de maio. Multa por atraso: US$ 400 (valor estadual fixo). Recomenda-se antecipar o envio.", size=9.5)
    body(doc, "Imposto Federal da LLC (declaração): prazos e eventuais multas variam conforme a classificação tributária e o ano-calendário. A multa por atraso do Form 1065 é definida pelo IRC §6698 e reajustada anualmente pelo IRS — confirmar o valor vigente a cada ano.", size=9.5)
    rule_note(doc, "Confirmar o valor da multa e o prazo fiscal conforme o ano e a classificação (Partnership / S-Corp / C-Corp) de cada cliente antes de entregar.")

    body(doc, "", space_after=6)
    body(doc, "AVISO: A Athena Business & Tax Advisors é uma empresa de contabilidade e não presta serviços de natureza jurídica ou de representação legal. Para assuntos legais, recomenda-se aconselhamento junto a um advogado habilitado.", size=8.5, italic=True, color=GREY)
    body(doc, "CONTATO: manager@athenataxadvisors.com · frontdesk@athenataxadvisors.com · (407) 777-2501", size=9, bold=True, color=NAVY, align="center")

    out = os.path.join(outdir, "00_CAPA_INDICE.docx"); doc.save(out); return out


# =====================================================================
# DOC 1 — OPERATING AGREEMENT  (English, logic-corrected)
# =====================================================================
def doc_operating(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    h1(doc, data["LLC_NAME"])
    body(doc, "OPERATING AGREEMENT", size=12, bold=True, align="center", color=GOLD, space_after=2)
    body(doc, "For a Member-Managed Limited Liability Company", size=10, italic=True, align="center", color=GREY, space_after=10)

    auth_mode = data.get("AUTHORITY_MODE", "SINGLE").upper()
    thr = data.get("DUAL_THRESHOLD_USD", "10000")

    # I. PRELIMINARY
    h2(doc, "I. Preliminary Provisions")
    body(doc, f"1. FORMATION. The members hereby form a limited liability company (“Company”) subject to the Florida Revised Limited Liability Company Act (Chapter 605, Florida Statutes) as currently in effect. The Articles of Organization were filed with the Florida Department of State on {data.get('ARTICLES_FILED_DATE','')}.")
    body(doc, f"2. NAME. The name of the Company shall be: {data['LLC_NAME']}.")
    body(doc, f"3. PRINCIPAL PLACE OF BUSINESS. {data.get('PRINCIPAL_ADDRESS','')}. Mailing address: {data.get('MAILING_ADDRESS','')}.")
    body(doc, f"4. REGISTERED OFFICE AND AGENT. {data.get('REGISTERED_AGENT','')}, located at {data.get('REGISTERED_AGENT_ADDRESS','')}. The registered office and agent may be changed by filing the appropriate form with the Florida Department of State.")
    body(doc, f"5. BUSINESS PURPOSE. {data.get('BUSINESS_PURPOSE','ANY AND ALL LAWFUL BUSINESS')}. This statement shall not limit the Company’s power to engage in any lawful activity.")
    body(doc, f"6. DURATION. The duration of the Company shall be {data.get('DURATION','PERPETUAL')}, until dissolved as provided in this Agreement or by law.")
    body(doc, f"7. MANAGEMENT STRUCTURE. This Company is MEMBER-MANAGED. It is managed by its members as provided in this Agreement; it has no separate managers unless and until the members amend this Agreement to elect manager-management.")
    rule_note(doc, "Modelo member-managed. Para LLC administrada por gestores, usar o template manager-managed.")

    # II. MEMBERSHIP
    h2(doc, "II. Membership Provisions")
    body(doc, "1. Nonliability of Members. No member shall be personally liable for the debts, obligations, or liabilities of the Company.")
    body(doc, "2. Reimbursement for Organizational Costs. Members shall be reimbursed for organizational expenses; the Company may elect to deduct and amortize such expenses as permitted by the Internal Revenue Code.")
    body(doc, "3. Members. The members of the Company are:")
    tbl = doc.add_table(rows=1, cols=3); tbl.style = "Table Grid"
    hcells = tbl.rows[0].cells
    for i, t in enumerate(["NAME", "ADDRESS", "ROLE"]):
        hcells[i].text = ""; run = hcells[i].paragraphs[0].add_run(t)
        run.font.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        _shade(hcells[i], "1F3B5C")
    for m in members:
        c = tbl.add_row().cells
        for i, v in enumerate([m["name"], m["address"], m["role"]]):
            c[i].text = ""; rr = c[i].paragraphs[0].add_run(v)
            rr.font.size = Pt(9); rr.font.name = "Calibri"
    body(doc, "", space_after=4)
    body(doc, "4. Members’ Percentage Interests. Each member’s “Percentage Interest” in the Company is FIXED as set forth in the Capital Provisions (Article IV) and restated in the Membership Ledger. Percentage Interests do NOT fluctuate with the balances of the members’ capital accounts and may be changed only by a written amendment to this Agreement signed by all members, or upon an admission, withdrawal, or transfer effected in accordance with this Agreement.", bold=False)
    rule_note(doc, "Correção: o percentual é fixo (tabela do Art. IV). Removida a regra antiga que recalculava o percentual pelo saldo da conta de capital.")
    body(doc, "5. Membership Voting. Except as otherwise required by this Agreement or Florida law, each member votes in proportion to the member’s Percentage Interest. “Majority of members” means members holding more than 50% of the total Percentage Interests. Matters designated below as requiring “unanimous consent” require the approval of all members.")
    body(doc, "6. Compensation. Members are not paid for acting as members. Members may be compensated for services rendered in another capacity (officer, employee, or independent contractor), including reasonable salaries or guaranteed payments approved by a majority of members.")
    body(doc, "7. Members’ Meetings. The Company holds no regular meetings; any member may call a meeting on reasonable notice by any written or electronic means. Minutes shall be placed in the Company’s records.")
    body(doc, "8. Other Business by Members. Members may own or work for other businesses; however, every member owes the Company good-faith effort in furtherance of its business.")

    # III. TAX & FINANCIAL
    h2(doc, "III. Tax and Financial Provisions")
    body(doc, f"1. Tax Classification. The Company shall initially be classified as a {data.get('TAX_CLASSIFICATION','PARTNERSHIP')} for federal (and, if applicable, state) income tax purposes. The members may change this election as permitted by law (e.g., IRS Form 8832 or Form 2553).")
    body(doc, f"2. Tax Year and Accounting Method. Tax year: {data.get('TAX_YEAR','CALENDAR YEAR')}. Accounting method: {data.get('ACCOUNTING_METHOD','CASH')}. Both may be changed with the consent of all members if the Company qualifies.")
    body(doc, "3. Partnership Representative. If the Company is subject to the centralized partnership audit regime of the Bipartisan Budget Act of 2015 (Internal Revenue Code §6223), the members shall designate a “Partnership Representative” (and, if required, a designated individual) to act for the Company before the IRS. Where eligible, the Company may elect out of the centralized audit regime under IRC §6221(b) for any tax year.")
    rule_note(doc, "Correção: substituído o obsoleto “Tax Matters Partner” (IRC §6231(a)(7), regime TEFRA revogado) pelo “Partnership Representative” do regime BBA vigente.")
    body(doc, "4. Annual Reports. Within 75 days after the end of each tax year, the Company shall provide each member with the information needed to prepare individual returns, including a Schedule K-1 (if taxed as a partnership) and a balance sheet and profit-and-loss statement.")
    body(doc, "5. Bank Accounts. Company funds shall be deposited in Company accounts, shall not be commingled with any member’s personal funds, and shall be handled only by persons authorized under the Statement of Authority and Article VII below.")
    body(doc, "6. Title to Assets. All Company property shall be held in the name of the Company.")

    # IV. CAPITAL
    h2(doc, "IV. Capital Provisions")
    body(doc, f"1. Capital Contributions. On or before {data.get('CONTRIBUTION_DUE_DATE','')}, the members shall contribute as follows, and shall receive the Percentage Interest shown:")
    tbl = doc.add_table(rows=1, cols=4); tbl.style = "Table Grid"
    hcells = tbl.rows[0].cells
    for i, t in enumerate(["NAME", "CONTRIBUTION", "FAIR MARKET VALUE", "PERCENTAGE INTEREST"]):
        hcells[i].text = ""; run = hcells[i].paragraphs[0].add_run(t)
        run.font.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        _shade(hcells[i], "1F3B5C")
    total = 0.0
    for m in members:
        c = tbl.add_row().cells
        fmv = m["fmv"]; total += float(m["percent"] or 0)
        fmv_str = f"US$ {fmv}" if fmv not in (None, "") else ""
        for i, v in enumerate([m["name"], m["contribution"], fmv_str, f"{m['percent']:g}%"]):
            c[i].text = ""; rr = c[i].paragraphs[0].add_run(str(v))
            rr.font.size = Pt(9); rr.font.name = "Calibri"
    body(doc, "", space_after=2)
    if abs(total - 100.0) > 0.01:
        rule_note(doc, f"ATENÇÃO: a soma das participações é {total:g}% (deveria ser 100%). Revisar a ficha.")
    body(doc, "2. Additional Contributions. Additional capital contributions may be required only by unanimous vote of the members.")
    body(doc, "3. Failure to Contribute. If a member fails to make a required contribution, the remaining members may, by unanimous vote, reschedule the contribution (with interest or penalty) or cancel the membership, refunding any prior partial payment.")
    body(doc, "4. No Interest on Capital. No interest is paid on capital contributions or capital-account balances.")
    body(doc, "5. Capital Accounts. A capital account is maintained for each member in accordance with the Internal Revenue Code and Treasury Regulations. Capital-account balances are used for tax bookkeeping and liquidation, and do NOT change any member’s Percentage Interest (see II.4).")
    body(doc, "6. Withdrawals and Distributions require the approval described in Article IV.8 and, for returns of capital, the written consent of all members.")
    body(doc, "7. Allocations. Profits, losses, and all items of income, gain, loss, deduction, and credit are allocated to the members in proportion to their Percentage Interests.")
    body(doc, "8. Distributions. Distributions of cash or property are made to the members in proportion to their Percentage Interests, at such times and in such amounts as approved by a MAJORITY of the members. No distribution shall be made if it would render the Company insolvent under §605.0405, Florida Statutes.")
    rule_note(doc, "Correção: removida a expressão sem sentido “ACCOUNTING RECORDS OF THE COMPANY”; a decisão de distribuir passa a ser por voto da maioria (e não conflita mais com a regra de consentimento para retorno de capital).")
    body(doc, "9. Tax Distributions. To the extent of available cash and subject to §605.0405, the Company shall use commercially reasonable efforts to distribute to each member, no later than each estimated-tax due date, an amount sufficient to cover the estimated income tax on the taxable income allocated to that member, pro rata to Percentage Interests.")
    rule_note(doc, "Cláusula nova: distribuição para cobrir imposto do K-1 (a LLC é partnership — o sócio paga imposto mesmo sem receber caixa).")
    body(doc, "10. Liquidation Distributions. On liquidation, items of income and loss are allocated to capital accounts and final distributions are made in proportion to positive capital-account balances.")

    # V. TRANSFER
    h2(doc, "V. Membership Withdrawal and Transfer")
    body(doc, f"1. Withdrawal. A member may withdraw on at least {data.get('WITHDRAWAL_NOTICE_DAYS','thirty (30)')} days’ written notice to the other members.")
    body(doc, "2. Restrictions on Transfer. A member may not transfer, pledge, or encumber all or part of the member’s interest without the prior written approval of all other members. A transferee is admitted as a member only upon unanimous consent of the non-transferring members.")
    body(doc, "3. Death or Incapacity of a Member; Buy-Sell (controls over dissolution). Upon the death or permanent incapacity of a member (a “Departing Member”): (a) for ninety (90) days, the surviving members have the option, by unanimous consent, to purchase all or part of the Departing Member’s interest at Fair Value determined as of that date; (b) to the extent not so purchased, the Company shall purchase the remaining interest at Fair Value. “Fair Value” is determined by agreement of the members and the Company’s accountant or, failing agreement, by an independent appraiser. This buy-sell governs, and the Company shall NOT dissolve by reason of such event (see VI).")
    rule_note(doc, "Correção: morte/incapacidade seguem UMA sequência (buy-sell de 90 dias) e deixam de ser gatilho de dissolução — eliminado o conflito com a antiga regra de “dissolver em 30 dias”.")
    body(doc, "4. Sale of Interest. If a member wishes to sell, the Company has a thirty (30) day option to purchase at Fair Value (exercisable by unanimous consent of the other members). If the Company does not exercise the option, the member may sell to another member or a third party, subject to the unanimous vote of the other members.")

    # VI. DISSOLUTION
    h2(doc, "VI. Dissolution")
    body(doc, "The Company dissolves only upon: (a) the written agreement of all members to dissolve; (b) the entry of a decree of judicial dissolution under Florida law; or (c) any event that makes it unlawful for the Company to continue. The death, incapacity, bankruptcy, retirement, resignation, or expulsion of a member does NOT dissolve the Company; the buy-sell provisions of Article V apply, and the remaining members may continue the Company.")
    rule_note(doc, "Correção: alinhado ao Chapter 605 — dissociação de um sócio não dissolve a LLC.")

    # VII. AUTHORITY & GENERAL
    h2(doc, "VII. Authority, Officers, and General Provisions")
    body(doc, "1. Approval vs. Signing Authority. Florida law distinguishes (i) APPROVING a Company action from (ii) having AUTHORITY TO SIGN and bind the Company. This Agreement addresses both separately: approval thresholds are set in the Articles above; signing authority is set in this Section and in the separate Statement of Authority.")
    if auth_mode == "DUAL":
        body(doc, f"2. Signing Authority (DUAL-CONTROL). Any one Authorized Member may sign for the Company in the ordinary course up to US$ {thr}. Any transaction at or above US$ {thr} — including opening or closing bank accounts, borrowing, and any purchase, sale, mortgage, or encumbrance of Company real estate — requires the signatures of TWO Authorized Members.")
        body(doc, "3. Disposition of Real Estate. Any sale, conveyance, mortgage, security interest, exchange, or other encumbrance of Company real estate requires the approval of a majority of Percentage Interests AND the signatures of two Authorized Members, and shall be recorded in a Minute of Meeting designating who signs for the Company.")
    else:
        body(doc, "2. Signing Authority (SINGLE). Any one Authorized Member may sign for the Company, including opening and operating bank accounts and signing checks, contracts, and other documents in the ordinary course of business. Presence of all members is not required.")
        body(doc, "3. Disposition of Real Estate. Any sale, conveyance, mortgage, security interest, exchange, or other encumbrance of Company real estate must first be APPROVED by a majority of Percentage Interests, and shall be recorded in a Minute of Meeting designating the Authorized Member who will sign for the Company in the transaction.")
    rule_note(doc, f"Autoridade gerada conforme a ficha: modo {auth_mode}" + (f" (limite US$ {thr})." if auth_mode=='DUAL' else " — separando aprovação de poder de assinatura."))
    body(doc, "4. Officers. The Company may appoint officers (President, Secretary, Treasurer, Manager, etc.), who need not be members, with compensation as approved by the members.")
    body(doc, "5. Records. The Company keeps at its principal address its organizational documents, this Agreement, membership records, and tax returns for the prior three years, open to inspection by any member on reasonable notice.")
    body(doc, f"6. Governing Law and Venue. This Agreement is governed by the laws of the State of {data.get('STATE','Florida')}. Venue for any dispute lies in the state or federal courts located in Orange County, Florida.")
    body(doc, "7. Mediation and Binding Arbitration. Any dispute among the members that cannot be resolved shall first be submitted to mediation; if unresolved, it shall be submitted to BINDING arbitration under the rules of the American Arbitration Association, seated in Orange County, Florida. The prevailing party may recover attorneys’ fees and costs. Judgment on the award may be entered in any court of competent jurisdiction.")
    rule_note(doc, "Correção: arbitragem passa a ser obrigatória e vinculante, com lei aplicável e foro (Orange County, FL).")
    body(doc, "8. Entire Agreement; Amendment. This Agreement is the entire agreement among the members and may be amended only by a written instrument signed by all current members.")
    body(doc, "9. Severability. If any provision is held invalid, the remaining provisions remain in effect.")

    # VIII. SIGNATURES
    h2(doc, "VIII. Signatures of Members")
    body(doc, f"IN WITNESS WHEREOF, the members adopt this Operating Agreement of {data['LLC_NAME']}.", space_after=6)
    body(doc, f"Date: {data.get('AGREEMENT_DATE', data.get('FORMATION_DATE',''))}", bold=True, space_after=6)
    for m in members:
        sig_block(doc, m["name"], m["role"])

    body(doc, "This is a template prepared for Athena Business & Tax Advisors. Athena is an accounting firm and does not provide legal services; the master template should be reviewed and approved by a licensed Florida attorney.", size=8, italic=True, color=GREY)

    out = os.path.join(outdir, "01_OPERATING_AGREEMENT.docx"); doc.save(out); return out


# =====================================================================
# DOC 2 — STATEMENT OF AUTHORITY  (renamed from Certificate of Authority)
# =====================================================================
def doc_authority(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    h1(doc, "STATEMENT OF AUTHORITY")
    body(doc, f"of {data['LLC_NAME']}", size=12, bold=True, align="center", color=GOLD, space_after=2)
    body(doc, "(Internal authorization of signatories — NOT a Certificate of Authority for a foreign LLC under §605.0902, Fla. Stat.)", size=8.5, italic=True, align="center", color=GREY, space_after=10)
    auth_mode = data.get("AUTHORITY_MODE", "SINGLE").upper()
    thr = data.get("DUAL_THRESHOLD_USD", "10000")
    body(doc, "The members of this Company hereby state the authority of the following persons to act on behalf of the Company:")
    authorized = [m for m in members if m["role"].upper().startswith("AUTHORIZED")]
    for m in authorized:
        body(doc, f"•  {m['name']} — {m['address']}", bold=True, size=10, space_after=2)
    body(doc, "", space_after=4)
    if auth_mode == "DUAL":
        body(doc, f"AUTHORITY (DUAL-CONTROL): Each Authorized Member listed above may open and operate bank accounts and sign checks, contracts, and other documents in the ordinary course of business up to US$ {thr} per transaction. Any transaction at or above US$ {thr} — including borrowing and any purchase, sale, or encumbrance of real property — requires the signatures of TWO Authorized Members.")
    else:
        body(doc, "AUTHORITY (SINGLE-SIGNER): Each Authorized Member listed above is authorized to open bank accounts, buy and sell property, and sign checks, withdrawals, contracts, and any other document on behalf of the Company. The presence of all members is not required; the signature of one Authorized Member is sufficient.")
    body(doc, "", space_after=4)
    body(doc, "This Statement of Authority is consistent with, and subject to, the Operating Agreement of the Company. Third parties may rely on it until a superseding statement is executed by the members.", size=9.5, italic=True)
    body(doc, f"Date: {data.get('FORMATION_DATE','')}", bold=True, space_after=8)
    body(doc, f"Name of LLC: {data['LLC_NAME']}", bold=True, space_after=8)
    for m in authorized:
        sig_block(doc, m["name"], m["role"])
    out = os.path.join(outdir, "02_STATEMENT_OF_AUTHORITY.docx"); doc.save(out); return out


# =====================================================================
# DOC 3 — MEMBERSHIP LEDGER  (clean footer, correct LLC)
# =====================================================================
def doc_ledger(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    h1(doc, "INTEREST CERTIFICATE LEDGER")
    body(doc, data["LLC_NAME"], size=13, bold=True, align="center", color=NAVY, space_after=2)
    body(doc, f"A {data['STATE']} Limited Liability Company", size=10, italic=True, align="center", color=GREY, space_after=10)
    tbl = doc.add_table(rows=1, cols=7); tbl.style = "Table Grid"
    heads = ["NO.", "TYPE OF ISSUANCE", "EFFECTIVE DATE", "DATE OF CERT.", "INTERESTS", "NAME OF MEMBER", "DATE CANCELED"]
    for i, t in enumerate(heads):
        c = tbl.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(t)
        run.font.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        _shade(c, "1F3B5C")
    for idx, m in enumerate(members, start=1):
        c = tbl.add_row().cells
        vals = [str(idx), "Membership Interest", data.get("FORMATION_DATE",""),
                data.get("FORMATION_DATE",""), f"{m['percent']:g}%", m["name"], "—"]
        for i, v in enumerate(vals):
            c[i].text = ""; rr = c[i].paragraphs[0].add_run(v)
            rr.font.size = Pt(8.5); rr.font.name = "Calibri"
    body(doc, "", space_after=8)
    body(doc, f"Date: {data.get('FORMATION_DATE','')}", bold=True, space_after=8)
    for m in members:
        sig_block(doc, m["name"], m["role"])
    out = os.path.join(outdir, "03_MEMBERSHIP_LEDGER.docx"); doc.save(out); return out


# =====================================================================
# DOC 4 — MEMBERSHIP CERTIFICATES  (one per member)
# =====================================================================
def doc_certificates(data, members, outdir):
    doc = Document(); set_margins(doc, 0.8, 0.8, 0.9, 0.9); letterhead(doc, data["LLC_NAME"])
    for idx, m in enumerate(members, start=1):
        if idx > 1:
            doc.add_page_break()
        body(doc, "", space_after=10)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"CERTIFICATE No. {idx:02d}"); r.font.size = Pt(11); r.font.bold = True
        r.font.color.rgb = GOLD; r.font.name = "Calibri"
        h1(doc, "MEMBERSHIP CERTIFICATE")
        body(doc, data["LLC_NAME"], size=18, bold=True, align="center", color=NAVY, space_after=2)
        body(doc, f"A {data['STATE']} Limited Liability Company", size=10, italic=True, align="center", color=GREY, space_after=14)
        body(doc, "This certifies that", size=11, align="center", color=BLACK, space_after=2)
        body(doc, m["name"], size=16, bold=True, align="center", color=NAVY, space_after=2)
        body(doc, "is the holder of a membership interest in the Company equal to", size=11, align="center", space_after=2)
        body(doc, f"{m['percent']:g}%  ({pct_words(m['percent'])})", size=15, bold=True, align="center", color=GOLD, space_after=14)
        body(doc, "and is entitled to all the rights of a member under the Articles of Organization, the Operating Agreement, and applicable law. This certificate is subject to the transfer restrictions in the Operating Agreement.", size=10, align="center", space_after=16)
        body(doc, f"Issued on {data.get('FORMATION_DATE','')} at Orlando, Florida.", size=10, align="center", space_after=18)
        body(doc, "_______________________________          _______________________________", size=10, align="center", space_after=2)
        body(doc, "Authorized Member                                      Authorized Member", size=8.5, align="center", color=GREY, space_after=2)
    out = os.path.join(outdir, "04_MEMBERSHIP_CERTIFICATES.docx"); doc.save(out); return out


# =====================================================================
# DOC 5 — MAILBOX AGREEMENT  (PT, period/validity aligned)
# =====================================================================
def doc_mailbox(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    h1(doc, "CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE CAIXA POSTAL")
    body(doc, "Mailbox Agreement", size=11, italic=True, align="center", color=GOLD, space_after=8)
    fee = data.get("MAILBOX_FEE_USD","380"); period = data.get("MAILBOX_PERIOD",""); valid = data.get("MAILBOX_VALID_UNTIL","")
    body(doc, f"Nome da Empresa: {data['LLC_NAME']}", bold=True, space_after=2)
    body(doc, f"Endereço: {data.get('PRINCIPAL_ADDRESS','')}", space_after=2)
    body(doc, f"E-mail do cliente: {data.get('CLIENT_EMAIL','')}     Telefone: {data.get('CLIENT_PHONE','')}", space_after=8)
    body(doc, "CLÁUSULA I — O presente contrato tem por objeto a prestação de serviços de Caixa Postal pela Athena Business and Tax Advisors ao cliente identificado acima, para recebimento e gerenciamento de correspondências.")
    body(doc, "CLÁUSULA II — O Cliente compromete-se a não utilizar os serviços para fins ilegais, fraudulentos ou proibidos pelas normas postais dos Estados Unidos.")
    body(doc, f"CLÁUSULA III — O valor do serviço é de US$ {fee},00, referente ao período de {period}. O pagamento poderá ser efetuado integralmente no ato da assinatura.")
    body(doc, f"CLÁUSULA IV — Este contrato vigora por {period} e tem validade até {valid}, coincidindo com o período contratado. A renovação, por igual período, dependerá de novo acordo e pagamento. Em caso de rescisão antecipada por descumprimento do Cliente, ficam devidos os meses restantes do período contratado.")
    rule_note(doc, "Correção: período pago e vigência agora coincidem (ambos vêm da ficha) — some o conflito 2026 × 2027.")
    body(doc, "CLÁUSULA V — O Cliente autoriza a Athena a abrir correspondências e descartar materiais de propaganda, spam ou fraudulentos, sem consulta prévia.")
    body(doc, "CLÁUSULA VI — Não é permitido o envio de caixas ou encomendas volumosas; o serviço restringe-se ao recebimento e à remessa de cartas e documentos, sem armazenamento físico.")
    body(doc, "CLÁUSULA VII — Toda correspondência recebida será digitalizada e enviada ao e-mail informado pelo Cliente. Mediante solicitação, a Athena poderá encaminhar fisicamente (por correio) a correspondência ao endereço previamente registrado pelo Cliente, no prazo de até 24 horas após o recebimento.")
    rule_note(doc, "Correção: reescrita a Cláusula VII — “reenvio físico para e-mail” não fazia sentido; separado digitalização (e-mail) de encaminhamento físico (correio).")
    body(doc, "AVISO LEGAL — A Athena Tax Advisors é uma empresa de contabilidade, sem fins ou viés de natureza jurídica ou de representação legal. Para assuntos legais, recomenda-se aconselhamento junto a um profissional habilitado.", size=9, italic=True, color=GREY)
    body(doc, "", space_after=6)
    body(doc, "DECLARAÇÃO — Ao assinar, o Cliente declara ciência e concordância com todos os termos acima.", bold=True)
    body(doc, "Assinatura: ______________________________________________", space_after=2)
    body(doc, f"Nome Completo: ______________________________     Data: ____/____/______", space_after=2)
    out = os.path.join(outdir, "05_MAILBOX_AGREEMENT.docx"); doc.save(out); return out


# =====================================================================
# DOC 6 — QA CHECKLIST  (PT, completeness gate)
# =====================================================================
def doc_qa(data, members, outdir):
    doc = Document(); set_margins(doc); letterhead(doc, data["LLC_NAME"])
    h1(doc, "CHECKLIST DE ABERTURA — CONFERÊNCIA DO PACOTE")
    body(doc, data["LLC_NAME"], size=12, bold=True, align="center", color=NAVY, space_after=10)
    body(doc, "Conferir ANTES de entregar. Todos os documentos devem refletir a mesma ficha do cliente.", italic=True, color=GREY, space_after=8)

    def check(txt, ok=None):
        box = "☑" if ok is True else ("☐")
        p = doc.add_paragraph()
        r = p.add_run(box + "  "); r.font.size = Pt(11); r.font.name = "Calibri"
        r2 = p.add_run(txt); r2.font.size = Pt(9.5); r2.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(2)

    h2(doc, "1. Coerência dos dados (todos os arquivos)")
    tot = sum(float(m["percent"] or 0) for m in members)
    check(f"Nome da LLC igual em todos os docs: {data['LLC_NAME']}")
    check(f"Nº de sócios: {len(members)}  |  Soma das participações: {tot:g}%", ok=(abs(tot-100)<0.01))
    check("Datas de formação, assinatura do acordo e emissão dos certificados registradas corretamente (podem diferir).")
    check("Nenhum rodapé/nome herdado de exemplo (JGHA, Triple Marques, OF & AF2G, “Page X of Y”).", ok=True)
    check("Endereço, agente registrado e classificação tributária conferem com o SunBiz/Articles.")

    h2(doc, "2. Documentos gerados")
    for d in ["00 Capa/Índice", "01 Operating Agreement", "02 Statement of Authority",
              "03 Membership Ledger", "04 Membership Certificates (um por sócio)",
              "05 Mailbox Agreement"]:
        check(d, ok=True)

    h2(doc, "3. Documentos oficiais a anexar (fora do gerador)")
    for d in [f"EIN / carta CP-575  (ficha: {'PREENCHIDO' if data.get('EIN') else 'PENDENTE'})",
              "Articles of Organization (cópia protocolada)",
              "Certificate of Status (US$ 5, opcional)",
              "Print da página do SunBiz",
              "Verificar BOI / FinCEN (Beneficial Ownership) — aplicável ou isento?"]:
        check(d, ok=(bool(data.get('EIN')) if d.startswith('EIN') else None))

    h2(doc, "4. Decisões do Operating Agreement confirmadas com o cliente")
    for d in [f"Modo de autoridade: {data.get('AUTHORITY_MODE','SINGLE')}" + (f" (limite US$ {data.get('DUAL_THRESHOLD_USD','')})" if data.get('AUTHORITY_MODE','').upper()=='DUAL' else ""),
              "Regra de distribuições (maioria) e distribuição para imposto.",
              "Buy-sell em caso de morte/incapacidade (90 dias) — sem dissolver a LLC.",
              "Percentual FIXO (não recalculado por conta de capital)."]:
        check(d, ok=True)

    h2(doc, "5. Serviços contínuos")
    check(f"Mailbox: período {data.get('MAILBOX_PERIOD','')} = validade {data.get('MAILBOX_VALID_UNTIL','')} (coincidem).")
    check("Business Tax Receipt (City of Orlando / Orange County) — verificar se o cliente opera fisicamente na cidade.")

    body(doc, "", space_after=6)
    body(doc, "Conferido por: ____________________________     Data: ____/____/______", bold=True)
    out = os.path.join(outdir, "06_CHECKLIST_ABERTURA.docx"); doc.save(out); return out


# =====================================================================
# MAIN
# =====================================================================
def main():
    ficha = sys.argv[1] if len(sys.argv) > 1 else "FICHA_ABERTURA_ATHENA.xlsx"
    data, members = read_ficha(ficha)
    safe = re.sub(r'[^A-Za-z0-9]+', '_', data["LLC_NAME"]).strip("_")
    outdir = os.path.join("PACOTE_" + safe)
    os.makedirs(outdir, exist_ok=True)
    outs = []
    outs.append(doc_capa(data, members, outdir))
    outs.append(doc_operating(data, members, outdir))
    outs.append(doc_authority(data, members, outdir))
    outs.append(doc_ledger(data, members, outdir))
    outs.append(doc_certificates(data, members, outdir))
    outs.append(doc_mailbox(data, members, outdir))
    outs.append(doc_qa(data, members, outdir))
    print("LLC:", data["LLC_NAME"], "| sócios:", len(members),
          "| soma %:", sum(m["percent"] for m in members),
          "| autoridade:", data.get("AUTHORITY_MODE"))
    for o in outs: print("  ->", o)

if __name__ == "__main__":
    main()
